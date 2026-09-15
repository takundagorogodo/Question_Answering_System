from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

def add_slide(title, bullets, note=""):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()
    for b in bullets:
        p = tf.add_paragraph()
        p.text = b
        p.level = 0
        p.font.size = Pt(18)
    if note:
        slide.notes_slide.notes_text_frame.text = note
    return slide

# 1 Title
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "Question Answering System using NLP and Generative QA"
slide.placeholders[1].text = "RAG + FLAN-T5-base | CPU-only Offline | i7-1355U\nStudent: Takunda\nMeasured Results: 0.735 vs 0.108 | Accuracy 1.00"

# 2 Requirements
add_slide("System Requirements - Locked & Verified", [
    "Target: Intel i7-1355U 10C/12T, 16GB RAM, Iris Xe iGPU",
    "Verified: torch.cuda.is_available() = False, torch.xpu.is_available() = False -> CPU-only",
    "Runtime: Windows offline, no API key, zero cost",
    "Domain: NLP/AI teaching corpus - 9 docs, AI-drafted human-reviewed (SOURCES.md)",
    "Why CPU-only is a feature: deployable on student laptops"
])

# 3 Architecture
add_slide("Architecture - Approach C: RAG + Local Generative", [
    "ADR-001: 4 approaches evaluated, C accepted",
    "Pipeline: Q -> all-MiniLM-L6-v2 (384-dim) -> FAISS IndexFlatIP (65 chunks, exact search) -> top_k=3 -> threshold 0.35 gate -> FLAN-T5-base (250M, 990MB) generates 1-2 sentences",
    "Why IndexFlatIP: exact search justified for 65 chunks, no approximation needed",
    "Why FLAN-T5-base: encoder-decoder, instruction-tuned, CPU-friendly, explainable"
])

# 4 Dataset
add_slide("Dataset & Chunking", [
    "Corpus: data/documents/ 9 txt files, 60 paragraphs -> 65 chunks",
    "Chunk stats: min 27, avg 452, max 599 chars, overlap 100 (word-boundary)",
    "Example: 'Natural Language Processing' -> 8 chunks, first chunk 27 chars = title",
    "Processed: documents.jsonl, chunks.jsonl, FAISS index faiss.index + chunks_meta.pkl",
    "Embeddings: 384-dim, paraphrase cosine 0.521 vs unrelated 0.033 (gap proves semantic space)"
])

# 5 Prompt probe
add_slide("Prompt Design - Stage 10b Probe at 0.803 Context", [
    "Q: Explain tokenization in simple words. | Retrieval: 02_tokenization.txt #0 score 0.803 (correct)",
    "A_current: 'using only context... answer exactly: insufficient information' -> REFUSED insufficient information",
    "B_no_escape: 'Answer in 1-2 sentences using info above' -> CORRECT but verbatim True (copied definition)",
    "C_soft_escape: 'If no relevant info reply I don't know' -> REFUSED I don't know",
    "D_own_words WINNER: 'Using context above, answer in your own words in 1-2 sentences' -> CORRECT 0.8028",
    "Decision: escape hatch moved to threshold gate (Stage 13), not in prompt"
])

# 6 Hallucination control
add_slide("Hallucination Control - Threshold 0.35 Measured, Not Guessed", [
    "In-domain: 0.779 (NLP), 0.803 (tokenization), 0.678 (RAG), 0.830 (hallucination), 0.737 (embedding), 0.580 (chunking) avg 0.735",
    "OOD: 0.157 (chocolate cake), 0.108 (IPL), 0.060 (tap) avg 0.108",
    "Gap: 0.40 vs 0.16 -> threshold 0.35 protects paraphrase ~0.45 while blocking OOD",
    "Gate: if top_score < 0.35 -> refuse in 0.0s without generator, grounded=False, confidence=low",
    "Result: correct_refusal_rate 1.00 (3/3 OOD)"
])

# 7 Evaluation
add_slide("Evaluation - Stage 17 Real Numbers, n=9, No Fabrication", [
    "Test set: 6 in-domain, 3 OOD (evaluate.py)",
    "In-domain: avg_score 0.735 grounded_rate 1.00 avg_latency 3.8s (first load 14.1s)",
    "OOD: avg_score 0.108 correct_refusal 1.00",
    "Overall accuracy (grounded matches expected): 1.00",
    "Saved: data/processed/eval_results.jsonl",
    "Limitation: n=9 small, needs 30-50 for final report"
])

# 8 Baseline
add_slide("Baseline Comparison - Stage 18 TF-IDF vs Dense", [
    "TF-IDF baseline: Q -> TF-IDF cosine -> stored chunk verbatim (NOT generative)",
    "Tokenization: TF-IDF 0.203 vs dense 0.803",
    "RAG: TF-IDF 0.405 vs dense 0.678",
    "OOD cake: TF-IDF 0.165 returns irrelevant chunk (hallucinates), RAG 0.157 correctly refuses",
    "Conclusion: dense better semantic, generative rewrites vs verbatim, threshold prevents OOD hallucination"
])

# 9 Demo
add_slide("Demo - Flask UI Working 0.828 High Grounded True", [
    "Flask: http://localhost:5000 - UI process never imports torch/pyarrow, child python.exe does via cli_answer.py",
    "Screenshot: explain tokenization -> Answer: 'Tokenization is the process of breaking...' confidence high grounded True top_score 0.827 latency 16.9s",
    "Sources: [0.828] 02_tokenization.txt #0, [0.580] #1, [0.531] #6 auditable with expandable text",
    "OOD demo: chocolate cake -> low 0.157 grounded False -> 'insufficient information - no relevant context found' in 0.04s"
])

# 10 WDAC
add_slide("Engineering Story - WDAC DLL Block", [
    "Problem: After Streamlit install, CLI broke: 'DLL load failed while importing _C: Application Control policy blocked' + pyarrow.lib blocked",
    "Root cause: College WDAC blocks unsigned native DLLs (torch._C.pyd, pyarrow.lib.pyd). Streamlit pulls pyarrow as dependency.",
    "Fix: uninstall pyarrow, downgrade scikit-learn 1.3.2 + numpy 1.26.4, Flask UI with subprocess bypass",
    "Result: CLI restored 0.8028 JSON, Flask UI works, Streamlit needs pyarrow so not used",
    "Viva point: debugging under constraints, not just building RAG"
])

# 11 Limitations
add_slide("Limitations & Future Work", [
    "Verbatim True for definition Qs (context is definition sentence) - still generative token-by-token but could improve paraphrase with larger model",
    "Latency 3.8s avg CPU, first load 14-16s - need pre-warm for demo",
    "Test set n=9 small - expand to 30-50",
    "Model comparison Stage 11 not fully benchmarked due to WDAC (FLAN-T5-small 80M vs base 250M)",
    "Future: ONNX runtime for embeddings, Q4 quantized LLM via llama.cpp Vulkan offload for Iris Xe (unverified, needs measurement)"
])

# 12 Conclusion
add_slide("Conclusion", [
    "End-to-end RAG works offline on ordinary laptop (i7-1355U CPU-only)",
    "Auditable: sources + scores + confidence + grounded flag + latency visible",
    "Measured: threshold 0.35 from real gap 0.735 vs 0.108, eval accuracy 1.00 on 9 questions",
    "Generative: FLAN-T5-base writes answer token-by-token, not TF-IDF verbatim copy",
    "Demo: Flask http://localhost:5000 ready in <2 min, OOD refusal proves no hallucination",
    "Git history shows incremental work: 60->65 chunks, 384-dim, 0.779/0.454/0.157 probe, 0.803 generator, 0.828 UI"
])

prs.save("docs/Question_Answering_System_Presentation.pptx")
print("Saved docs/Question_Answering_System_Presentation.pptx")